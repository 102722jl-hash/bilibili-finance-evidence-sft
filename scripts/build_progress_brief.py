from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"D:\桌面\workspace\宏观2\项目最新进展汇报.docx"
BLUE = "2E74B5"
NAVY = "0B2545"
GRAY = "555555"
LIGHT = "F2F4F7"
GREEN = "1F6B45"
GOLD = "7A5A00"
RED = "9B1C1C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E0E8", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=11, color="000000", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(p, before=0, after=6, line=1.1):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_para(doc, text="", size=11, color="000000", bold=False, italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph()
    style_paragraph(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        set_run_font(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=13, after=6)
        set_run_font(p.add_run(text), size=15, color=BLUE, bold=True)
    else:
        style_paragraph(p, before=8, after=4)
        set_run_font(p.add_run(text), size=12, color=NAVY, bold=True)
    return p


def add_bullet(doc, text, color="000000"):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=3, line=1.08)
    set_run_font(p.add_run(text), size=10.5, color=color)
    return p


def add_status_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [1800, 2520, 2520, 2520])
    set_table_borders(table)
    headers = ["指标", "当前结果", "状态", "说明"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        set_run_font(p.add_run(value), size=10, color=NAVY, bold=True)
    rows = [
        ("正式案例", "10 个", "已完成", "10 个独立事件全部进入 V3 正式集"),
        ("规范 QA", "90 道", "已完成", "每案 9 道，覆盖六类用户意图"),
        ("对话路径", "50 条", "已完成", "2-6 轮各 10 条，支持连续追问"),
        ("SFT 训练候选", "72 + 18 条", "进行中", "72 条训练，18 条独立验证"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            style_paragraph(p, after=0)
            color = GREEN if value == "已完成" else GOLD if value == "进行中" else "000000"
            set_run_font(p.add_run(value), size=10, color=color, bold=value in {"已完成", "进行中"})
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(header, after=0)
    set_run_font(header.add_run("证据链数据项目 · 阶段性进展"), size=8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(footer, after=0)
    set_run_font(footer.add_run("2026年8月12日"), size=8.5, color=GRAY)

    add_para(doc, "项目最新进展汇报", size=22, color=NAVY, bold=True, after=2)
    add_para(doc, "财经视频事件推理链数据集建设", size=12, color=GRAY, after=10)
    meta = doc.add_table(rows=1, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(meta, [4680, 4680])
    set_table_borders(meta, color="FFFFFF", size="0")
    for cell in meta.rows[0].cells:
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=100, bottom=100)
    p = meta.cell(0, 0).paragraphs[0]
    style_paragraph(p, after=0)
    set_run_font(p.add_run("当前阶段："), size=9.5, color=NAVY, bold=True)
    set_run_font(p.add_run("研究数据正式版已完成"), size=9.5, color=GREEN, bold=True)
    p = meta.cell(0, 1).paragraphs[0]
    style_paragraph(p, after=0)
    set_run_font(p.add_run("下一阶段："), size=9.5, color=NAVY, bold=True)
    set_run_font(p.add_run("生成并验证 SFT 训练文件"), size=9.5, color=GOLD, bold=True)

    add_heading(doc, "一、总体结论", 1)
    add_para(doc, "项目已完成从视频/文章材料到“案例事实图谱 + QA 评测接口”的正式 V3 数据层建设。当前核心研究资产已经具备可审计、可复核、可重复构建的条件；可直接用于模型训练的 SFT 文件正在由该数据层派生，尚未作为最终训练交付物宣布完成。", after=6)

    add_heading(doc, "二、已完成的核心数据成果", 1)
    add_status_table(doc)
    add_para(doc, "10 个案例均包含三条完整链：事实传导链、作者推理链、外部核验链。数据保存步骤、证据引用、支持状态和不确定性边界，不保存模型自由式思维过程。", size=10.2, color=GRAY, before=4, after=4)

    add_heading(doc, "三、数据质量与证据边界", 1)
    add_bullet(doc, "严格区分 author_trace（作者如何推导）与 world_evidence（现实证据支持到哪里），作者推理忠实不等于现实因果成立。")
    add_bullet(doc, "90 道 QA 全部为 answerable/complete；“可回答”表示能够完整说明证据边界，不表示问题中的现实命题已经被证实。")
    add_bullet(doc, "没有直接反证时不标记 contradicted；未获支持的命题保留 unsupported 或 not_reviewed。")
    add_bullet(doc, "粮食、日本、保健品三个案例已从候选状态恢复到正式集；保健品案例的赞助披露保持 unknown。")
    add_bullet(doc, "正式 CLI 校验为 valid；完整仓库测试此前为 224 passed；V3 发布文件重复构建哈希一致，V2 原始数据哈希不变。")

    doc.add_page_break()
    add_heading(doc, "四、当前 SFT 训练集进展", 1)
    add_para(doc, "目标是生成通用三消息 messages JSONL，供主流监督微调框架直接读取：", after=3)
    code = add_para(doc, '{"messages": [{"role": "system"}, {"role": "user"}, {"role": "assistant"}]}', size=9.5, color=NAVY, after=6)
    code.paragraph_format.left_indent = Inches(0.18)
    add_bullet(doc, "system：规定证据边界、引用要求和拒答规则。")
    add_bullet(doc, "user：提供短材料、问题、真值范围和候选引用。")
    add_bullet(doc, "assistant：输出直接答案、三条审计链、引用、可回答性和不确定性说明。")
    add_para(doc, "当前 SFT 阶段状态：单条样本生成器已经实现，计划输出 72 条训练样本和 18 条独立验证样本；发布器测试和 CLI 接入仍在进行，因此 data/train/v3 暂不作为最终交付目录。", after=7)

    add_heading(doc, "五、当前风险与限制", 1)
    add_bullet(doc, "90 条 QA 适合验证 SFT 管线和小规模领域适配，不足以单独训练具备广泛泛化能力的基础模型。")
    add_bullet(doc, "部分作者归因仍缺乏现实外证，训练时必须保留证据边界，不能把作者观点学习成无条件事实。")
    add_bullet(doc, "训练集必须与验证集按案例和事件簇隔离，避免同一事件的改写样本造成评估污染。")

    add_heading(doc, "六、下一步工作", 1)
    for text in [
        "完成 V3 SFT 发布器测试和数据校验。",
        "接入 train-build / train-validate 的 V3 路由。",
        "生成 data/train/v3：72 条 train、18 条 validation。",
        "重复构建并比较哈希，完成完整回归测试。",
        "向训练团队交付 SFT 文件、数据卡和使用说明。",
    ]:
        p = doc.add_paragraph(style="List Number")
        style_paragraph(p, after=3, line=1.08)
        set_run_font(p.add_run(text), size=10.5)

    add_heading(doc, "七、汇报口径", 1)
    add_para(doc, "可以对外汇报为：研究数据正式版已完成，形成 10 个事件案例、90 道高价值 QA 和 50 条连续追问路径；SFT 训练文件正在由已验证数据层派生，预计完成后再进入模型训练。", color=NAVY, bold=True, before=2, after=0)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
