from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def main():
    path = Path(sys.argv[1])
    doc = Document(path)
    assert len(doc.sections) == 1
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert round(section.left_margin.inches, 2) == 1.0
    assert round(section.right_margin.inches, 2) == 1.0

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    required = [
        "财经视频事件推理链数据集\n阶段成果汇报",
        "一、完成成果与当前状态",
        "二、已经形成的研究数据链路",
        "三、质量控制与证据边界",
        "四、QA 问答设计的核心思考",
        "五、阶段交付物",
    ]
    for text in required:
        assert text in paragraphs, text

    all_text = "\n".join(paragraphs)
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + " | ".join(cell.text for cell in row.cells)
    for token in ("10 个", "90 条", "50 条", "valid: true", "QA 问答设计的核心思考"):
        assert token in all_text, token
    for forbidden in ("待填写", "TODO", "占位符", "SFT 派生文件", "SFT 训练数据", "72 train"):
        assert forbidden not in all_text, forbidden

    expected_grids = [
        [9360],
        [1900, 1260, 1700, 4500],
        [9360],
        [2300, 2200, 4860],
        [9360],
        [2500, 4680, 2180],
    ]
    assert len(doc.tables) == len(expected_grids)
    for table, expected in zip(doc.tables, expected_grids):
        actual = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert actual == expected, (actual, expected)
        for row in table.rows:
            for cell, width in zip(row.cells, expected):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert tc_w is not None and int(tc_w.get(qn("w:w"))) == width

    assert section.header.paragraphs[0].text == "FinGoose 财经事件研究数据 | 阶段成果汇报"
    assert section.footer.paragraphs[0].text == "内部阶段材料 | 截至 2026-08-12"

    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        names = set(package.namelist())
        for part in ("word/document.xml", "word/styles.xml", "word/header1.xml", "word/footer1.xml"):
            assert part in names

    print(
        {
            "valid": True,
            "paragraphs": len(paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "file_size": path.stat().st_size,
        }
    )


if __name__ == "__main__":
    main()
